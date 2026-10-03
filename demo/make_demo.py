# -*- coding: utf-8 -*-
"""生成示例课件，方便第一次运行就能看到效果。

会生成两份文件到 demo/ 目录：
    计算机网络_第3章_传输层.pptx
    操作系统_第2章_进程管理.pdf

这两份课件是故意做得不那么干净的：里面混了页脚、页码、重复页眉、
结尾致谢页，用来验证清洗模块真的有用。

用法：
    python demo/make_demo.py

依赖：python-pptx（生成 PPT）、reportlab（生成 PDF）。
reportlab 只在生成示例时需要，正常使用本工具不需要装它。
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfgen import canvas

DEMO_DIR = Path(__file__).resolve().parent

# PPT 里反复出现的页眉，用来测试重复文字会被当成页眉去掉
PPT_HEADER = "计算机网络 · 第 3 章 传输层"


def add_page_number(slide, page_no: int) -> None:
    """在右下角加一个页码，这是最典型的装饰文字。"""
    box = slide.shapes.add_textbox(Inches(8.6), Inches(6.9), Inches(1.0), Inches(0.4))
    box.text_frame.text = str(page_no)


def add_header(slide) -> None:
    """在顶部加一条重复页眉。"""
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.15), Inches(8.0), Inches(0.4))
    box.text_frame.text = PPT_HEADER


def add_content_slide(prs: Presentation, title: str, bullets: list[str], page_no: int):
    """加一页「标题 + 正文」的幻灯片。"""
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title

    body = slide.placeholders[1].text_frame
    body.text = bullets[0]
    for line in bullets[1:]:
        paragraph = body.add_paragraph()
        paragraph.text = line

    for paragraph in body.paragraphs:
        paragraph.font.size = Pt(20)

    add_header(slide)
    add_page_number(slide, page_no)
    return slide


def build_pptx(path: Path) -> None:
    """生成示例 PPT 课件。"""
    prs = Presentation()
    page_no = 1

    # 封面
    cover = prs.slides.add_slide(prs.slide_layouts[0])
    cover.shapes.title.text = "计算机网络 · 第 3 章 传输层"
    cover.placeholders[1].text_frame.text = "期末复习课件"
    add_page_number(cover, page_no)

    slides = [
        (
            "3.1 传输层概述",
            [
                "传输层是指为两台主机中的应用进程提供端到端通信服务的层次。",
                "端到端通信是指数据从发送方进程直接送到接收方进程，中间网络不拆开看。",
                "本章重点：端口、UDP、TCP、拥塞控制。",
            ],
        ),
        (
            "3.2 端口与套接字",
            [
                "端口是指传输层用来标识应用进程的 16 位编号。",
                "端口号范围是 0 ~ 65535。",
                "套接字是指 IP 地址与端口号的组合。",
                "注意区分：端口标识进程，IP 地址标识主机。",
            ],
        ),
        (
            "3.3 UDP 协议",
            [
                "UDP 是一种无连接的传输层协议。",
                "UDP 的特点是开销小、时延低，但不保证可靠交付。",
                "常见错误：把 UDP 当成可靠传输协议，这是典型易错点。",
                "因此 UDP 适合实时音视频这类能容忍少量丢包的场景。",
            ],
        ),
        (
            "3.4 TCP 协议",
            [
                "TCP 是一种面向连接的、提供可靠交付的传输层协议。",
                "TCP 通过序号、确认、重传三项机制实现可靠传输。",
                "TCP 建立连接需要三次握手，释放连接需要四次挥手。",
                "为什么 TCP 建立连接需要三次握手？",
            ],
        ),
        (
            "3.5 传输效率计算",
            [
                "吞吐量 = 窗口大小 / 往返时间 RTT",
                "信道利用率 U = 吞吐量 / 带宽",
                "结论：窗口越大吞吐量越高，但受网络拥塞限制。",
            ],
        ),
        (
            "3.6 拥塞控制",
            [
                "拥塞控制是指防止过多数据注入网络、导致网络性能下降的机制。",
                "慢开始阶段，拥塞窗口从 1 开始按指数增长。",
                "到达慢开始门限后进入拥塞避免，改为线性增长。",
                "结论：因此拥塞控制是站在网络整体角度考虑问题。",
            ],
        ),
        (
            "3.7 易错点提醒",
            [
                "易错点：流量控制是端到端的，拥塞控制面向整个网络，两者不要混淆。",
                "注意区分：流量控制靠接收窗口，拥塞控制靠拥塞窗口。",
                "了解内容：TCP 选项字段的扩展用法属于选学，考试一般不考。",
            ],
        ),
        (
            "本章思考题",
            [
                "简述 TCP 与 UDP 的主要区别。",
                "试述 TCP 三次握手的完整过程。",
                "说明滑动窗口协议的作用与实现步骤。",
            ],
        ),
        (
            "本章小结",
            [
                "传输层的核心是端到端通信、端口寻址与可靠传输。",
                "重点掌握：端到端通信、端口、UDP、TCP、拥塞控制。",
            ],
        ),
    ]

    for title, bullets in slides:
        page_no += 1
        add_content_slide(prs, title, bullets, page_no)

    # 结尾页，专门用来测试致谢页会被清洗掉
    page_no += 1
    end = prs.slides.add_slide(prs.slide_layouts[5])
    end.shapes.title.text = "谢谢观看"
    add_page_number(end, page_no)

    prs.save(str(path))


def build_pdf(path: Path) -> None:
    """生成示例 PDF 课件（用 reportlab 画出来）。"""
    # 注册中文字体，否则 PDF 里的中文会显示成方块
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))

    pdf = canvas.Canvas(str(path), pagesize=A4)
    width, height = A4

    pages = [
        (
            "操作系统 · 第 2 章 进程管理",
            [
                "进程是指程序在一个数据集合上运行的过程，是系统进行资源分配的基本单位。",
                "线程是指进程内部的一条执行流，是 CPU 调度的基本单位。",
                "进程控制块 PCB 是指用来描述进程状态和控制信息的数据结构。",
            ],
        ),
        (
            "2.2 进程状态与转换",
            [
                "进程的三种基本状态是就绪、运行和阻塞。",
                "注意：就绪状态和阻塞状态的区别在于是否只差 CPU。",
                "易错点：把等待输入输出当成就绪状态，是常见错误。",
                "结论：因此阻塞态必须先回到就绪态，不能直接进入运行态。",
            ],
        ),
        (
            "本章思考题",
            [
                "简述进程与线程的区别。",
                "说明进程三种基本状态之间的转换条件。",
            ],
        ),
    ]

    for page_no, (title, lines) in enumerate(pages, start=1):
        pdf.setFont("STSong-Light", 16)
        pdf.drawString(60, height - 80, title)

        pdf.setFont("STSong-Light", 12)
        y = height - 130
        for line in lines:
            pdf.drawString(60, y, line)
            y -= 28

        # 页脚，同样用来测试清洗效果
        pdf.setFont("STSong-Light", 9)
        pdf.drawString(60, 50, f"第 {page_no} 页")
        pdf.showPage()

    pdf.save()


def main() -> None:
    """生成两份示例课件。"""
    pptx_path = DEMO_DIR / "计算机网络_第3章_传输层.pptx"
    pdf_path = DEMO_DIR / "操作系统_第2章_进程管理.pdf"

    build_pptx(pptx_path)
    build_pdf(pdf_path)

    print(f"已生成：{pptx_path}")
    print(f"已生成：{pdf_path}")


if __name__ == "__main__":
    main()
